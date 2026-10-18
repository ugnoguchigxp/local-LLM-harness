"""Offline, single-process CPU probes; artifacts are written outside Git."""
import argparse
import base64
import csv
from datetime import datetime, timezone
import fcntl
import hashlib
import io
import json
import math
import os
from pathlib import Path
import platform
import resource
import socket
import statistics
import subprocess
import sys
import time
import traceback

sys.dont_write_bytecode = True


def cer(expected, actual):
    """Character edit distance; whitespace is intentionally retained."""
    previous = list(range(len(actual) + 1))
    for row, a in enumerate(expected, 1):
        current = [row]
        for col, b in enumerate(actual, 1):
            current.append(min(current[-1] + 1, previous[col] + 1, previous[col-1] + (a != b)))
        previous = current
    return {'edits': previous[-1], 'reference_characters': len(expected),
            'cer': previous[-1] / len(expected) if expected else None}


def nms(boxes, axis):
    kept = []
    for item in sorted(boxes, key=lambda item: item['score'], reverse=True):
        lo, hi = item['box'][axis], item['box'][axis+2]
        if any(max(0, min(hi, b['box'][axis+2])-max(lo, b['box'][axis])) /
               max(1, max(hi, b['box'][axis+2])-min(lo, b['box'][axis])) > .5 for b in kept):
            continue
        kept.append(item)
    return sorted(kept, key=lambda b: b['box'][axis])


def table_grid(objects, tokens):
    """Small explicit adapter for the pilot; raw model output is always retained."""
    rows = nms([o for o in objects if o['label'] == 'table row'], 1)
    cols = nms([o for o in objects if o['label'] == 'table column'], 0)
    cells, spans = [], []
    for span in [o for o in objects if o['label'] == 'table spanning cell']:
        x1, y1, x2, y2 = span['box']
        rr = [i for i, r in enumerate(rows) if y1 <= (r['box'][1]+r['box'][3])/2 <= y2]
        cc = [i for i, c in enumerate(cols) if x1 <= (c['box'][0]+c['box'][2])/2 <= x2]
        if rr and cc and len(rr)*len(cc) > 1:
            spans.append({'rows': rr, 'columns': cc})
    for ri, row in enumerate(rows):
        for ci, col in enumerate(cols):
            box = [col['box'][0], row['box'][1], col['box'][2], row['box'][3]]
            assigned = []
            for token in tokens:
                b = token['box']
                x, y = (b[0]+b[2])/2, (b[1]+b[3])/2
                if box[0] <= x < box[2] and box[1] <= y < box[3]:
                    assigned.append(token)
            assigned.sort(key=lambda t: (round(t['box'][1]/12), t['box'][0]))
            cells.append({'row': ri, 'column': ci, 'box': box,
                          'text': ' '.join(t['text'] for t in assigned)})
    return {'rows': len(rows), 'columns': len(cols), 'cells': cells, 'merged': spans,
            'adapter': 'pilot-center-assignment; not full upstream canonicalization'}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--stage', choices=['baseline', 'florence', 'table', 'docling', 'd1'], required=True)
    parser.add_argument('--evidence', type=Path, required=True)
    parser.add_argument('--run-id', required=True)
    parser.add_argument('--warm-repeats', type=int, default=20)
    parser.add_argument('--smoke-only', action='store_true')
    parser.add_argument('--florence-implementation', choices=['native', 'upstream'], default='upstream')
    parser.add_argument('--baseline-result', default='baseline-main.json')
    parser.add_argument('--baseline-python-source', type=Path, default=Path('/srv/ai/apps/vision-ocr-host'))
    parser.add_argument('--baseline-models', type=Path, default=Path('/srv/ai/local-services-prepared/vision-ocr-production/models/models.json'))
    args = parser.parse_args()
    if args.evidence.resolve().is_relative_to(Path(__file__).resolve().parents[2]):
        raise ValueError('evidence must stay outside the source repository')
    if not args.run_id.replace('-', '').replace('_', '').isalnum():
        raise ValueError('invalid run id')
    args.evidence.mkdir(parents=True, exist_ok=True)
    inference_lock = (args.evidence / '.inference.lock').open('a')
    fcntl.flock(inference_lock, fcntl.LOCK_EX)
    output = args.evidence / (args.stage + '-' + args.run_id + '.json')
    if output.exists():
        raise ValueError('run already exists; select a new run id')
    os.environ['HF_HUB_OFFLINE'] = '1'
    os.environ['TRANSFORMERS_OFFLINE'] = '1'
    os.environ['HF_DATASETS_OFFLINE'] = '1'
    os.environ['TOKENIZERS_PARALLELISM'] = 'false'

    def no_network(*a, **kw):
        raise RuntimeError('network forbidden during inference')
    socket.create_connection = no_network
    socket.socket.connect = no_network
    fixtures_manifest = args.evidence / 'fixtures' / 'manifest.json'
    fixtures = {item['id']: item for item in json.loads(fixtures_manifest.read_text())['fixtures']}
    for fixture in fixtures.values():
        path = Path(fixture['path'])
        if hashlib.sha256(path.read_bytes()).hexdigest() != fixture['sha256']:
            raise ValueError('fixture checksum mismatch')
    result = {'stage': args.stage, 'run_id': args.run_id, 'device': 'cpu', 'threads': 4,
              'started_utc': datetime.now(timezone.utc).isoformat(), 'executable': sys.executable,
              'arguments': sys.argv[1:],
              'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              'source_revision': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=Path(__file__).resolve().parents[2], text=True).strip(),
              'host_loadavg': Path('/proc/loadavg').read_text().strip(),
              'python': platform.python_version(), 'platform': platform.platform(),
              'offline': True, 'smoke_only': args.smoke_only, 'warm_repeats': args.warm_repeats,
              'fixtures_sha256': hashlib.sha256(fixtures_manifest.read_bytes()).hexdigest(),
              'criteria': 'exploratory task-level evidence; production thresholds not set',
              'observations': [], 'errors': [], 'complete': False}

    def write():
        result['peak_process_rss_mib'] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')

    def observe(fixture_id, task, function):
        started = time.perf_counter()
        try:
            value = function()
            observation = {'fixture': fixture_id, 'task': task, 'output': value,
                           'elapsed_ms': (time.perf_counter()-started)*1000}
            result['observations'].append(observation)
            write()
            print(json.dumps({'fixture': fixture_id, 'task': task,
                              'elapsed_ms': round(observation['elapsed_ms'], 2)}, ensure_ascii=False), flush=True)
            return value
        except Exception as error:
            result['errors'].append({'fixture': fixture_id, 'task': task,
                                     'error': repr(error), 'traceback': traceback.format_exc()})
            write()
            print('FAILED ' + fixture_id + ' ' + task + ': ' + repr(error), flush=True)
            return None

    def benchmark(function):
        durations = []
        for _ in range(args.warm_repeats):
            started = time.perf_counter()
            function()
            durations.append((time.perf_counter()-started)*1000)
        if durations:
            values = sorted(durations)
            result['warm'] = {'n': len(values), 'p50_ms': statistics.median(values),
                              'p95_ms': values[math.ceil(.95*len(values))-1],
                              'max_ms': values[-1], 'all_ms': durations}
        write()

    try:
        from PIL import Image
        if args.stage == 'baseline':
            sys.path.insert(0, str(args.baseline_python_source))
            from vision_ocr.engines import OcrEngine, ObjectEngine, load_image
            import cv2
            config = {'models': json.loads(args.baseline_models.read_text())}
            result['models'] = config['models']
            result['opencv'] = cv2.__version__
            started = time.perf_counter()
            ocr = OcrEngine(config)
            objects = ObjectEngine(config)
            result['load_ms'] = (time.perf_counter()-started)*1000
            ids = ['japanese-text'] if args.smoke_only else ['japanese-text', 'blank', 'table-en', 'table-ja', 'document-mixed-ja']
            for fixture_id in ids:
                value = observe(fixture_id, 'ocr', lambda fid=fixture_id: ocr.infer(load_image(fixtures[fid]['path'])))
                if value and 'text_lines' in fixtures[fixture_id]['expected']:
                    value['reference_comparison'] = cer('\n'.join(fixtures[fixture_id]['expected']['text_lines']), value['text'])
                    write()
            if not args.smoke_only:
                for fixture_id in ['bus', 'car', 'blank']:
                    observe(fixture_id, 'objects', lambda fid=fixture_id: objects.infer(load_image(fixtures[fid]['path'])))
            benchmark(lambda: ocr.infer(load_image(fixtures['japanese-text']['path'])))
        elif args.stage == 'd1':
            sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'apps' / 'd1-system-one'))
            from engine import Engine
            started = time.perf_counter()
            engine = Engine('/srv/ai/models/d1-3b', device='cpu')
            result['load_and_warmup_ms'] = (time.perf_counter()-started)*1000
            result['engine'] = engine.health()

            def predict(fixture_id):
                image = Image.open(fixtures[fixture_id]['path']).convert('RGB')
                image.thumbnail((384, 384))
                stream = io.BytesIO()
                image.save(stream, format='PNG')
                return engine.predict({'model': 'd1-3b-game', 'state': 'Judge only the visible image.',
                    'images': ['data:image/png;base64,' + base64.b64encode(stream.getvalue()).decode()],
                    'questions': {
                        'person': {'type': 'noul', 'instructions': 'Is a person visibly present in the image?'},
                        'table': {'type': 'noul', 'instructions': 'Does the image contain a structured data table with rows and columns?'},
                        'formula': {'type': 'noul', 'instructions': 'Does the image contain a mathematical equation?'}}})
            for fixture_id in (['bus'] if args.smoke_only else ['bus', 'car', 'blank', 'table-en', 'table-ja', 'formula-pythagoras', 'document-mixed-ja']):
                observe(fixture_id, 'three-questions-384px', lambda fid=fixture_id: predict(fid))
            benchmark(lambda: predict('bus'))
        else:
            import torch
            import transformers
            torch.set_num_threads(4)
            torch.set_num_interop_threads(1)
            result['torch'] = torch.__version__
            result['transformers'] = transformers.__version__
            result['dtype'] = 'float32'
            model_records = json.loads((args.evidence / 'models.json').read_text())

            def snapshot(name):
                record = model_records[name]
                root = Path(record['directory'])
                for entry in record['files']:
                    with (root / entry['path']).open('rb') as stream:
                        digest = hashlib.file_digest(stream, 'sha256').hexdigest()
                    if digest != entry['sha256']:
                        raise ValueError('snapshot checksum mismatch')
                result.setdefault('models', {})[name] = record
                return root

            if args.stage == 'florence':
                from transformers import AutoProcessor
                root = snapshot('florence')
                started = time.perf_counter()
                result['implementation'] = args.florence_implementation
                if args.florence_implementation == 'upstream':
                    from transformers import AutoModelForCausalLM
                    processor = AutoProcessor.from_pretrained(root, local_files_only=True, trust_remote_code=True)
                    model = AutoModelForCausalLM.from_pretrained(root, torch_dtype=torch.float32,
                        local_files_only=True, trust_remote_code=True, attn_implementation='eager').eval()
                else:
                    from transformers import Florence2ForConditionalGeneration
                    processor = AutoProcessor.from_pretrained(root, local_files_only=True, trust_remote_code=False)
                    model = Florence2ForConditionalGeneration.from_pretrained(root, dtype=torch.float32,
                        local_files_only=True, attn_implementation='sdpa').eval()
                result['load_ms'] = (time.perf_counter()-started)*1000
                result['generation'] = {'max_new_tokens': 128, 'num_beams': 3, 'do_sample': False}

                def predict(fixture_id, task, text=''):
                    image = Image.open(fixtures[fixture_id]['path']).convert('RGB')
                    inputs = processor(text=task+text, images=image, return_tensors='pt')
                    with torch.inference_mode():
                        ids = model.generate(**inputs, **result['generation'])
                    raw = processor.batch_decode(ids, skip_special_tokens=False)[0]
                    parsed = processor.post_process_generation(raw, task=task, image_size=image.size)
                    return {'raw': raw, 'parsed': parsed, 'tokens': ids.shape[-1],
                            'at_token_limit': ids.shape[-1] >= 128}
                ids = ['car'] if args.smoke_only else ['car', 'bus', 'blank', 'japanese-text', 'table-ja']
                for fixture_id in ids:
                    observe(fixture_id, '<CAPTION>', lambda fid=fixture_id: predict(fid, '<CAPTION>'))
                if not args.smoke_only:
                    for task in ['<DETAILED_CAPTION>', '<DENSE_REGION_CAPTION>', '<OD>']:
                        for fixture_id in ['car', 'bus']:
                            observe(fixture_id, task, lambda fid=fixture_id, t=task: predict(fid, t))
                    observe('car', '<CAPTION_TO_PHRASE_GROUNDING>', lambda: predict('car', '<CAPTION_TO_PHRASE_GROUNDING>', 'A car.'))
                benchmark(lambda: predict('car', '<CAPTION>'))
            elif args.stage == 'table':
                from transformers import AutoConfig, AutoImageProcessor, AutoModelForObjectDetection
                root = snapshot('table')
                det_root = snapshot('table-detection')
                started = time.perf_counter()
                processor = AutoImageProcessor.from_pretrained(root, local_files_only=True)
                config = AutoConfig.from_pretrained(root, local_files_only=True)
                config.use_pretrained_backbone = False
                det_config = AutoConfig.from_pretrained(det_root, local_files_only=True)
                det_config.use_pretrained_backbone = False
                result['backbone_init'] = 'checkpoint only; no separate pretrained-backbone download'
                model, loading_info = AutoModelForObjectDetection.from_pretrained(root, config=config,
                    local_files_only=True, output_loading_info=True)
                detector, det_loading_info = AutoModelForObjectDetection.from_pretrained(det_root, config=det_config,
                    local_files_only=True, output_loading_info=True)
                model.eval()
                detector.eval()
                result['loading_info'] = {
                    name: {key: sorted(value) if isinstance(value, set) else value for key, value in info.items()}
                    for name, info in [('structure', loading_info), ('detection', det_loading_info)]}
                for info in [loading_info, det_loading_info]:
                    if info.get('missing_keys') or info.get('mismatched_keys'):
                        raise ValueError('missing or mismatched checkpoint tensors')
                    if any(not key.endswith('num_batches_tracked') for key in info.get('unexpected_keys', [])):
                        raise ValueError('unexpected checkpoint tensors')
                det_processor = AutoImageProcessor.from_pretrained(det_root, local_files_only=True)
                result['load_ms'] = (time.perf_counter()-started)*1000
                result['threshold'] = .5

                def detect(fixture_id, structure=True):
                    image = Image.open(fixtures[fixture_id]['path']).convert('RGB')
                    active, prep = (model, processor) if structure else (detector, det_processor)
                    with torch.inference_mode():
                        value = active(**prep(images=image, return_tensors='pt'))
                    post = prep.post_process_object_detection(value, threshold=.5,
                        target_sizes=torch.tensor([[image.height, image.width]]))[0]
                    return [{'label': active.config.id2label[int(label)], 'score': float(score),
                             'box': [round(float(v), 2) for v in box]}
                            for label, score, box in zip(post['labels'], post['scores'], post['boxes'])]
                baseline_path = args.evidence / args.baseline_result
                baseline = json.loads(baseline_path.read_text()) if baseline_path.exists() else {'observations': []}
                for fixture_id in (['table-en'] if args.smoke_only else ['table-en', 'table-ja', 'table-merged', 'blank']):
                    objects = observe(fixture_id, 'structure', lambda fid=fixture_id: detect(fid))
                    if objects is None:
                        continue
                    expected = fixtures[fixture_id]['expected']
                    oracle = [{'text': cell['text'], 'box': cell['text_bbox']} for cell in expected.get('cells', [])]
                    oracle_grid = table_grid(objects, oracle)
                    result['observations'][-1]['oracle_grid'] = oracle_grid
                    baseline_ocr = next((item['output'] for item in baseline['observations'] if item['fixture'] == fixture_id and item['task'] == 'ocr'), {})
                    tokens = []
                    for line in baseline_ocr.get('lines', []):
                        points = line['polygon']
                        tokens.append({'text': line['text'], 'box': [min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)]})
                    if tokens:
                        grid = table_grid(objects, tokens)
                        result['observations'][-1]['ocr_grid'] = grid
                        rows = [['' for _ in range(grid['columns'])] for _ in range(grid['rows'])]
                        for cell in grid['cells']:
                            rows[cell['row']][cell['column']] = cell['text']
                        exports = args.evidence / ('exports-' + args.stage + '-' + args.run_id)
                        exports.mkdir(exist_ok=True)
                        with (exports / (fixture_id + '.csv')).open('w', newline='') as stream:
                            csv.writer(stream).writerows(rows)
                        from openpyxl import Workbook, load_workbook
                        workbook = Workbook()
                        for row in rows:
                            workbook.active.append(row)
                        for span in grid['merged']:
                            workbook.active.merge_cells(start_row=min(span['rows'])+1, end_row=max(span['rows'])+1,
                                start_column=min(span['columns'])+1, end_column=max(span['columns'])+1)
                        xlsx = exports / (fixture_id + '.xlsx')
                        workbook.save(xlsx)
                        reread = load_workbook(xlsx, read_only=False)
                        result['observations'][-1]['xlsx_reread'] = [list(row) for row in reread.active.values]
                    write()
                if not args.smoke_only:
                    for fixture_id in ['document-mixed-en', 'document-mixed-ja', 'blank']:
                        observe(fixture_id, 'table-detection', lambda fid=fixture_id: detect(fid, False))
                benchmark(lambda: detect('table-en'))
            elif args.stage == 'docling':
                from transformers import AutoProcessor, AutoModelForImageTextToText
                root = snapshot('docling')
                started = time.perf_counter()
                processor = AutoProcessor.from_pretrained(root, local_files_only=True)
                model = AutoModelForImageTextToText.from_pretrained(root, dtype=torch.float32,
                    local_files_only=True, attn_implementation='sdpa').eval()
                result['load_ms'] = (time.perf_counter()-started)*1000

                def predict(fixture_id, instruction):
                    image = Image.open(fixtures[fixture_id]['path']).convert('RGB')
                    messages = [{'role': 'user', 'content': [{'type': 'image'}, {'type': 'text', 'text': instruction}]}]
                    prompt = processor.apply_chat_template(messages, add_generation_prompt=True)
                    inputs = processor(text=prompt, images=[image], return_tensors='pt')
                    budget = 192 if instruction == 'Convert formula to LaTeX.' else 1024
                    with torch.inference_mode():
                        ids = model.generate(**inputs, max_new_tokens=budget, do_sample=False, use_cache=True)
                    generated = ids[:, inputs['input_ids'].shape[1]:]
                    raw = processor.batch_decode(generated, skip_special_tokens=False)[0]
                    return {'raw': raw, 'tokens': generated.shape[-1], 'max_new_tokens': budget,
                            'at_token_limit': generated.shape[-1] >= budget}
                for fixture_id in (['formula-pythagoras'] if args.smoke_only else ['formula-pythagoras', 'formula-quadratic', 'blank']):
                    observe(fixture_id, 'formula', lambda fid=fixture_id: predict(fid, 'Convert formula to LaTeX.'))
                if not args.smoke_only:
                    for fixture_id in ['table-en', 'table-ja', 'document-mixed-en', 'document-mixed-ja']:
                        observe(fixture_id, 'document', lambda fid=fixture_id: predict(fid, 'Convert this page to docling.'))
                benchmark(lambda: predict('formula-pythagoras', 'Convert formula to LaTeX.'))
        result['complete'] = True
    except Exception as error:
        result['errors'].append({'error': repr(error), 'traceback': traceback.format_exc()})
        print(traceback.format_exc(), flush=True)
    finally:
        result['finished_utc'] = datetime.now(timezone.utc).isoformat()
        write()
    if result['errors']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
