"""Purpose-specific offline image chains; no answer key is accepted by this API."""
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import time
import unicodedata

from evaluate import table_grid

POLICY = json.loads(Path(__file__).with_name('integration-policy.json').read_text())


def overlap_fraction(inner, outer):
    area = max(0, inner[2]-inner[0])*max(0, inner[3]-inner[1])
    intersection = max(0, min(inner[2], outer[2])-max(inner[0], outer[0])) * max(0, min(inner[3], outer[3])-max(inner[1], outer[1]))
    return intersection / area if area else 0


def validate_grid(grid, tokens, rule_grid=None):
    issues = []
    if not grid['rows'] or not grid['columns'] or not tokens:
        issues.append('no_cells_or_text')
    ambiguous, unmatched = 0, 0
    for token in tokens:
        matches = sorted((overlap_fraction(token['box'], c['box']) for c in grid['cells']), reverse=True)
        if not matches or matches[0] < POLICY['token_cell_overlap']:
            unmatched += 1
        elif len(matches) > 1 and matches[1] > .2:
            ambiguous += 1
    if unmatched:
        issues.append('text_crosses_or_misses_cells')
    if ambiguous:
        issues.append('ambiguous_text_assignment')
    if grid['merged']:
        issues.append('merged_cells_not_qualified')
    if rule_grid:
        if not rule_grid.get('rectangular', True):
            issues.append('nonrectangular_ruling_cells')
        if (grid['rows'], grid['columns']) != (rule_grid['rows'], rule_grid['columns']):
            issues.append('ruling_grid_shape_conflict')
        predicted = {(tuple(s['rows']), tuple(s['columns'])) for s in grid['merged']}
        ruled = {(tuple(s['rows']), tuple(s['columns'])) for s in rule_grid['merged']}
        if predicted != ruled:
            issues.append('ruling_merge_conflict')
    else:
        issues.append('independent_structure_check_unavailable')
    return {'issues': issues, 'unmatched_tokens': unmatched, 'ambiguous_tokens': ambiguous}


def export_grid(grid, destination):
    """Keep literal OCR strings, including leading zeros and '=' in XLSX."""
    from openpyxl import Workbook, load_workbook
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    rows = [['' for _ in range(grid['columns'])] for _ in range(grid['rows'])]
    for cell in grid['cells']:
        rows[cell['row']][cell['column']] = cell['text']
    with destination.with_suffix('.csv').open('w', newline='', encoding='utf-8') as stream:
        csv.writer(stream, quoting=csv.QUOTE_ALL).writerows(rows)
    book = Workbook()
    for ri, row in enumerate(rows, 1):
        for ci, text in enumerate(row, 1):
            cell = book.active.cell(ri, ci, text)
            cell.data_type = 's'
    for span in grid['merged']:
        book.active.merge_cells(start_row=min(span['rows'])+1, end_row=max(span['rows'])+1,
                               start_column=min(span['columns'])+1, end_column=max(span['columns'])+1)
    xlsx = destination.with_suffix('.xlsx')
    book.save(xlsx)
    loaded = load_workbook(xlsx, data_only=False)
    return {'csv': str(destination.with_suffix('.csv')), 'xlsx': str(xlsx),
            'reread': [[cell.value for cell in row] for row in loaded.active],
            'types': [[cell.data_type for cell in row] for row in loaded.active],
            'merged': [str(value) for value in loaded.active.merged_cells.ranges]}


def formula_atoms(text, latex=False):
    text = unicodedata.normalize('NFKC', text)
    if latex:
        text = re.sub(r'\\[A-Za-z]+', '', text)
    return tuple(re.findall(r'[A-Za-z]|[0-9]+|[+\-=]', text))


class ImageChains:
    def __init__(self, backend, output):
        self.backend = backend
        self.output = Path(output)
        if self.output.resolve().is_relative_to(Path(__file__).resolve().parents[2]):
            raise ValueError('artifacts must stay outside the source repository')
        self.output.mkdir(parents=True, exist_ok=True)

    def run(self, image_path, objective, request_id, input_kind='auto'):
        from PIL import Image, ImageOps, ImageStat
        if objective not in ('ocr', 'describe', 'table', 'formula', 'document', 'judge'):
            raise ValueError('unknown objective')
        if input_kind not in ('auto', 'photo', 'document'):
            raise ValueError('unknown input kind')
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', request_id):
            raise ValueError('invalid request id')
        directory = self.output / request_id
        directory.mkdir(exist_ok=False)
        started = time.perf_counter()
        result = {'request_id': request_id, 'objective': objective, 'input': str(image_path),
                  'status': 'failed', 'steps': [], 'issues': [], 'policy': POLICY['version'],
                  'schema_version': 'vision-chain-pilot.v1', 'input_kind': input_kind}

        def call(op, path):
            stamp = time.perf_counter()
            try:
                value = self.backend.call(op, path)
            except Exception as error:
                result['steps'].append({'operation': op, 'elapsed_ms': (time.perf_counter()-stamp)*1000,
                                        'status': 'failed', 'error': repr(error)})
                raise
            result['steps'].append({'operation': op, 'elapsed_ms': (time.perf_counter()-stamp)*1000,
                                    'model': value.get('model'), 'revision': value.get('revision'),
                                    'operation_first_use': value.get('operation_first_use')})
            return value

        try:
            image_path = Path(image_path)
            result['input_sha256'] = hashlib.sha256(image_path.read_bytes()).hexdigest()
            with Image.open(image_path) as source:
                if source.width*source.height > POLICY['max_pixels'] or max(source.size) > POLICY['max_side']:
                    raise ValueError('image_dimensions_exceeded')
                if source.format not in ('PNG', 'JPEG', 'WEBP'):
                    raise ValueError('unsupported_image')
                result['raw_size'] = list(source.size)
                result['exif_orientation'] = source.getexif().get(274, 1)
                image = ImageOps.exif_transpose(source).convert('RGB')
            result['size'] = list(image.size)
            result['coordinate_frame'] = 'pixels of EXIF-normalized input; origin top-left'
            canonical = directory / 'normalized.png'
            image.save(canonical)
            result['normalized_input'] = str(canonical)
            extrema = ImageStat.Stat(image).extrema
            result['channel_ranges'] = [high-low for low, high in extrema]
            if max(result['channel_ranges']) <= POLICY['blank_channel_range']:
                result.update(status='empty', reason='uniform_image', model_calls=0)
            elif objective == 'describe':
                objects = call('objects', canonical) if input_kind == 'auto' else {'objects': []}
                known = [o for o in objects['objects'] if o['confidence'] >= .4]
                if input_kind == 'photo' or known:
                    caption = call('caption', canonical)
                    result.update(status='needs_review' if caption['at_token_limit'] else 'ready',
                                  description=caption['text'], objects=known, raw=caption['raw'],
                                  description_kind='model_caption', factual_guarantee=False)
                else:
                    text = call('ocr', canonical)
                    if text['lines']:
                        result.update(status='needs_review', description_kind='extracted_text',
                                      description=text['text'], text=text,
                                      reason='scene_description_unavailable_text_extracted')
                    else:
                        result.update(status='needs_review', reason='no_supported_scene_or_text',
                                      objects=objects['objects'])
            elif objective == 'ocr':
                text = call('ocr', canonical)
                result.update(status='ready' if text['lines'] else 'not_detected', text=text)
            elif objective == 'table':
                self._tables(image, canonical, directory, result, call)
            elif objective == 'formula':
                text = call('ocr', canonical)
                formula = call('formula', canonical)
                result.update(formula=formula, ocr=text)
                if formula['at_token_limit']:
                    result.update(status='needs_review', issues=['generation_limit'])
                elif not formula.get('latex'):
                    result.update(status='needs_review', issues=['no_formula'])
                elif formula_atoms(text['text']) != formula_atoms(formula['latex'], latex=True):
                    result.update(status='needs_review', issues=['ocr_formula_symbol_disagreement'])
                else:
                    result.update(status='needs_review', issues=['formula_structure_requires_review'],
                                  verification='ordered visible symbols agree; mathematical structure is not independently verified')
            elif objective == 'document':
                text = call('ocr', canonical)
                document = call('document', canonical)
                result.update(document=document, ocr=text, status='needs_review')
                if document['at_token_limit']:
                    result['issues'].append('generation_limit')
                # Whole-page conversion has no independent structure proof in this pilot.
                result['issues'].append('document_structure_requires_review')
                source_text = re.sub(r'\s+', '', unicodedata.normalize('NFKC', text['text']))
                absent = [cell for cell in document.get('table_text', []) if cell and
                          re.sub(r'\s+', '', unicodedata.normalize('NFKC', cell)) not in source_text]
                if absent:
                    result['issues'].append('table_text_not_supported_by_ocr')
                    result['unsupported_table_text'] = absent
            else:
                result.update(status='ready', judgment=call('judge', canonical), factual_guarantee=False)
        except Exception as error:
            result.update(status='failed', error=repr(error))
        result['elapsed_ms'] = (time.perf_counter()-started)*1000
        result['operation_calls'] = len(result['steps'])
        result['model_calls'] = sum(step['operation'] != 'layout' for step in result['steps'])
        (directory / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
        return result

    def _tables(self, image, canonical, directory, result, call):
        layout = call('layout', canonical)
        detections = call('table-detection', canonical)
        boxes = [d['box'] for d in detections['objects'] if d['score'] >= POLICY['detection_threshold']]
        proposals = []
        # Use visible ruling geometry to retain complete table boundaries rather than text-only crops.
        for region in layout['regions']:
            proposals.append({'box': region['box'], 'rule_grid': region['grid'], 'region_source': 'opencv_ruling'})
        for box in boxes:
            if not any(overlap_fraction(box, item['box']) > .5 or overlap_fraction(item['box'], box) > .5 for item in proposals):
                proposals.append({'box': box, 'rule_grid': None, 'region_source': 'table_transformer_detection'})
        result['detection'] = detections
        result['layout'] = layout
        if not proposals:
            result.update(status='not_detected', reason='no_table_candidate', tables=[])
            return
        if len(proposals) > POLICY['max_tables']:
            result.update(status='needs_review', issues=['too_many_table_candidates'], tables=[])
            return
        tables = []
        for index, proposal in enumerate(proposals):
            box = proposal['box']
            crop_box = [max(0, math.floor(box[0])-12), max(0, math.floor(box[1])-12),
                        min(image.width, math.ceil(box[2])+12), min(image.height, math.ceil(box[3])+12)]
            if crop_box[2] <= crop_box[0] or crop_box[3] <= crop_box[1]:
                raise ValueError('invalid_table_crop')
            crop = directory / ('table-'+str(index)+'.png')
            image.crop(crop_box).save(crop)
            structure = call('table-structure', crop)
            text = call('ocr', crop)
            tokens = []
            for line in text['lines']:
                points = line['polygon']
                tokens.append({'text': line['text'], 'box': [min(p[0] for p in points), min(p[1] for p in points),
                                                           max(p[0] for p in points), max(p[1] for p in points)]})
            grid = table_grid(structure['objects'], tokens)
            model_grid = grid
            rules = proposal['rule_grid']
            geometry_refinement = None
            if rules and (grid['rows'], grid['columns']) == (rules['rows'], rules['columns']) and not rules['merged'] and not grid['merged']:
                xs = [x-crop_box[0] for x in rules['xs']]
                ys = [y-crop_box[1] for y in rules['ys']]
                aligned = [{'label': 'table row', 'score': 1, 'box': [xs[0], ys[i], xs[-1], ys[i+1]]} for i in range(len(ys)-1)]
                aligned.extend({'label': 'table column', 'score': 1, 'box': [xs[i], ys[0], xs[i+1], ys[-1]]} for i in range(len(xs)-1))
                grid = table_grid(aligned, tokens)
                geometry_refinement = 'visible ruling boundaries after model shape agreement'
            validation = validate_grid(grid, tokens, proposal['rule_grid'])
            original_cells = [{**cell, 'box': [cell['box'][0]+crop_box[0], cell['box'][1]+crop_box[1],
                                              cell['box'][2]+crop_box[0], cell['box'][3]+crop_box[1]]} for cell in grid['cells']]
            table = {'source_box': crop_box, 'region_source': proposal['region_source'], 'grid': grid,
                     'source_cells': original_cells, 'validation': validation, 'model_grid': model_grid,
                     'geometry_refinement': geometry_refinement,
                     'status': 'needs_review' if validation['issues'] else 'ready',
                     'objects': structure['objects'], 'ocr': text}
            table['exports'] = export_grid(grid, directory / ('table-'+str(index)))
            tables.append(table)
        result.update(tables=tables, status='ready' if all(t['status'] == 'ready' for t in tables) else 'needs_review')
